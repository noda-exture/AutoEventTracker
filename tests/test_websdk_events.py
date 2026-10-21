import copy
import unittest

from openpyxl import Workbook
from measurement_adapters.adobe_analytics import AdobeAnalyticsAdapter
from excel_reporter import load_project_config, process_single_sheet


class WebSDKEventsTests(unittest.TestCase):
    def setUp(self):
        self.adapter = AdobeAnalyticsAdapter()
        self.packet = {'type': 'AEP Web SDK', 'step_id': '0001', 'step_index': 0,
                       'xdm_payload': {'events': [{'xdm': {'_experience': {'analytics': {
                           'event1to100': {'event10': {'value': 1}, 'event2': {'value': 5},
                                          'event1': {'value': 1}, 'event3': {'value': 1, 'id': 'abc'},
                                          'event4': {'value': 0}},
                           'event101to200': {'event101': {'value': 2.5}}
                       }}}}]}}

    def test_xdm_values_ids_and_numeric_order_are_preserved(self):
        before = copy.deepcopy(self.packet)
        self.assertEqual(self.adapter.extract_value(self.packet, {'source_key': 'events'}),
                         'event1,event2=5,event3:abc,event4=0,event10,event101=2.5')
        self.assertEqual(self.packet, before)

    def test_explicit_data_events_including_empty_override_xdm(self):
        event = self.packet['xdm_payload']['events'][0]
        for value in ('event9,event8=3', ''):
            event['data'] = {'__adobe': {'analytics': {'events': value}}}
            self.assertEqual(self.adapter.extract_value(self.packet, {'source_key': 'events'}), value)

    def test_appmeasurement_is_unchanged(self):
        self.assertEqual(self.adapter.extract_value(
            {'type': 'Adobe Analytics (Legacy)', 'params': {'events': 'event9,event1=2'}},
            {'source_key': 'events'}), 'event9,event1=2')

    def test_excel_existing_events_row_and_raw_fields(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Adobe Analytics'
        ws.append(['Adobe Analytics'])
        ws.append(['送信文字列', 'XDM', 'Data', '名前', '説明'])
        ws.append(['events', '', 'data.__adobe.analytics.events', 'Events', ''])
        process_single_sheet(ws, ws.title, [self.packet], [], load_project_config('/nonexistent'),
                             'latest.json', '', compare=False)
        self.assertEqual(ws['F3'].value, 'event1,event2=5,event3:abc,event4=0,event10,event101=2.5')
        self.assertTrue(any('event2.value' in str(ws.cell(row, 2).value)
                            for row in range(4, ws.max_row + 1)))
        wb.close()

    def test_missing_events_row_is_discovered(self):
        rows = self.adapter.discover_mappings([self.packet], [])
        self.assertEqual(sum(row['source_key'] == 'events' for row in rows), 1)


if __name__ == '__main__':
    unittest.main()
