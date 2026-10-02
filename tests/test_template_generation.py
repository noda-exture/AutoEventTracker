import os
import tempfile
import unittest

from openpyxl import load_workbook

from create_template import create_comprehensive_template
from measurement_adapters.adobe_analytics import AdobeAnalyticsAdapter, get_nested_value
from measurement_adapters.ga4 import GA4Adapter
from excel_reporter import process_single_sheet, load_project_config


class TemplateGenerationTests(unittest.TestCase):
    def test_report_appends_unknown_keys_from_both_versions(self):
        with tempfile.TemporaryDirectory() as directory:
            path = create_comprehensive_template(os.path.join(directory, 'template.xlsx'))
            workbook = load_workbook(path)
            sheet = workbook['GA4']
            original_last = sheet.max_row
            def packet(params):
                return {'type': 'GA4', 'step_id': 'step_1', 'step_index': 0,
                        'params': {'en': 'page_view', 'tid': 'G-TEST', **params}}
            process_single_sheet(sheet, 'GA4',
                                 [packet({'ep.new_field': 'new', 'ep.shared': 'after'})],
                                 [packet({'up.old_field': 'old', 'ep.shared': 'before'})],
                                 load_project_config(directory), 'new.json', 'old.json')
            keys = [sheet.cell(row, 1).value for row in range(original_last + 1, original_last + 4)]
            self.assertEqual(keys, ['ep.new_field', 'ep.shared', 'up.old_field'])
            self.assertEqual(sheet.cell(original_last + 1, 6).value, 'new')
            self.assertEqual(sheet.cell(original_last + 2, 6).value, 'after')
            old_start = original_last + 3 + 4
            self.assertEqual(sheet.cell(old_start + original_last, 6).value, 'before')
            self.assertEqual(sheet.cell(old_start + original_last + 1, 6).value, 'old')
            workbook.save(os.path.join(directory, 'report.xlsx'))
            original = load_workbook(path)
            self.assertEqual(original['GA4'].max_row, original_last)
            original.close()
            workbook.close()

    def test_adobe_unknown_nested_fields_and_covered_subtrees(self):
        adapter = AdobeAnalyticsAdapter()
        existing = [dict(source_key='g', primary_path='web.webPageDetails.URL',
                         secondary_path='data.__adobe.analytics.pageURL', label='URL'),
                    dict(source_key='', primary_path='productListItems', secondary_path='', label='Items')]
        packet = {'params': {'custom': 'legacy'}, 'xdm_payload': {'events': [{
            'xdm': {'web': {'webPageDetails': {'URL': 'https://example.test'}},
                    '_tenant': {'custom': 0}, 'productListItems': [{'SKU': 'one'}]},
            'data': {'__adobe': {'analytics': {'g': 'https://example.test',
                    'contextData': {'custom.key': False}}}}
        }]}}
        discovered = adapter.discover_mappings([packet, packet], existing)
        self.assertEqual(len(discovered), 3)
        extracted = {m['label']: adapter.extract_value(packet, m) for m in discovered}
        self.assertEqual(extracted, {'custom': 'legacy', 'xdm._tenant.custom': 0,
                                    'data.__adobe.analytics.contextData.custom.key': False})

    def test_adobe_template_extracts_official_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            path = create_comprehensive_template(os.path.join(directory, 'template.xlsx'))
            workbook = load_workbook(path)
            mappings = {
                row[3]: dict(source_key=row[0] or '', primary_path=row[1] or '',
                             secondary_path=row[2] or '', label=row[3])
                for row in workbook['Adobe Analytics'].iter_rows(min_row=3, values_only=True)
            }
            adapter = AdobeAnalyticsAdapter()
            legacy = {'params': {'pe': 'lnk_d', 'pev2': 'Download PDF', 'events': 'event1'}}
            self.assertEqual(adapter.extract_value(legacy, mappings['Link Type']), 'lnk_d')
            self.assertEqual(adapter.extract_value(legacy, mappings['Link Name']), 'Download PDF')
            packet = {'xdm_payload': {'events': [{
                'xdm': {
                    'eventType': 'commerce.purchases',
                    'identityMap': {'ECID': [{'id': '123456789'}]},
                    'commerce': {'purchases': {'value': 1},
                                 'order': {'payments': [{'transactionID': 'T1'}]}},
                    'web': {'webPageDetails': {'URL': 'https://xdm.example/'},
                            'webInteraction': {'name': 'Download PDF', 'type': 'download'}},
                    '_experience': {'analytics': {'customDimensions': {
                        'props': {'prop75': 'last prop'}, 'eVars': {'eVar250': 'last evar'}}}},
                },
                'data': {'__adobe': {'analytics': {'g': ''}}},
            }]}}
            for label, expected in [('Link Name', 'Download PDF'), ('Link Type', 'download'),
                                    ('Experience Cloud ID', '123456789'), ('Transaction ID', 'T1'),
                                    ('XDM Purchases', 1), ('Events', None),
                                    ('XDM Event Type', 'commerce.purchases'), ('Page URL', ''),
                                    ('prop75', 'last prop'), ('eVar250', 'last evar')]:
                with self.subTest(label=label):
                    self.assertEqual(adapter.extract_value(packet, mappings[label]), expected)
            self.assertNotIn('eVar251', mappings)
            self.assertNotIn('prop76', mappings)
            self.assertIn('list3', mappings)
            self.assertFalse(any('_tenant' in row['primary_path'] for row in mappings.values()))
            workbook.close()

    def test_array_paths_preserve_missing_zero_and_literal_keys(self):
        self.assertEqual(get_nested_value({'items': [{'value': 0}]}, 'items[0].value'), 0)
        self.assertIsNone(get_nested_value({'items': []}, 'items[0].value'))
        self.assertIsNone(get_nested_value({'items': {}}, 'items[0].value'))
        self.assertEqual(get_nested_value({'items[0]': 'literal'}, 'items[0]'), 'literal')


if __name__ == '__main__':
    unittest.main()
