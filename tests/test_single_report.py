import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from PySide6.QtWidgets import QApplication
from openpyxl import Workbook, load_workbook

from display_parts import Ui_MainWindow
from excel_reporter import generate_compare_report
from main import AutoTrackerApp


class SingleReportTests(unittest.TestCase):
    def test_single_report_has_values_without_comparison_rows_or_rules(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            wb = Workbook()
            sheet = wb.active
            sheet.title = 'GA4'
            sheet.append([])
            sheet.append(['GA4パラメータ', '種別', '型', '項目', '説明'])
            sheet.append(['en', '', '', 'イベント名', ''])
            wb.save(project / 'template.xlsx')
            data = {'events': [{'type': 'GA4', 'step_index': 0, 'step_id': '0001',
                                'params': {'en': 'page_view', 'tid': 'G-TEST'}}]}
            latest = project / 'latest.json'
            latest.write_text(json.dumps(data))
            result = generate_compare_report(str(project), str(latest), None, 'single.xlsx', logger_func=lambda _: None)
            report = load_workbook(result)
            sheet = report['GA4']
            self.assertEqual(sheet['F3'].value, 'page_view')
            self.assertEqual(len(sheet.conditional_formatting), 0)
            text = '\n'.join(str(cell.value or '') for row in sheet for cell in row)
            for marker in ('過去データ', '比較元になし', '最新のみ', '比較元のみ'):
                self.assertNotIn(marker, text)
            report.close()

    def test_recording_report_preserves_receive_order_and_rejects_comparison(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            wb = Workbook()
            wb.active.title = 'GA4'
            wb.active.append([])
            wb.active.append(['GA4パラメータ', '種別', '型', '項目', '説明'])
            wb.active.append(['en', '', '', 'イベント名', ''])
            wb.save(project / 'template.xlsx')
            data = {'capture_source': 'recording', 'events': [
                {'type': 'GA4', 'capture_sequence': 1, 'timestamp': 2,
                 'url': 'https://www.google-analytics.com/g/collect?tid=G-T',
                 'post_data': 'en=first\nen=second'},
                {'type': 'GA4', 'capture_sequence': 2, 'timestamp': 1,
                 'params': {'en': 'third'}}]}
            latest = project / 'recording.json'
            latest.write_text(json.dumps(data))
            result = generate_compare_report(str(project), str(latest), None, 'recording.xlsx')
            report = load_workbook(result)
            sheet = report['GA4']
            self.assertEqual([sheet.cell(3, col).value for col in (6, 7, 8)],
                             ['first', 'second', 'third'])
            self.assertTrue(sheet['F2'].value.startswith('通信 1'))
            self.assertTrue(sheet['H2'].value.startswith('通信 2'))
            self.assertNotIn('Step', sheet['F2'].value)
            self.assertNotIn('実行結果', report.sheetnames)
            report.close()
            logs = []
            generate_compare_report(str(project), str(latest), str(latest), 'compare.xlsx', logs.append)
            self.assertIn('比較をオフ', logs[-1])
            self.assertFalse((project / 'outputs' / 'compare.xlsx').exists())

    def test_unconfirmed_hover_is_reported_even_without_packets(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            wb = Workbook()
            wb.active.title = 'GA4'
            wb.save(project / 'template.xlsx')
            data = {'events': [], 'execution_status': 'partial', 'step_results': [
                {'step_number': 1, 'step_id': '0001', 'action': 'hover',
                 'selector': '=literal', 'status': 'unconfirmed'}]}
            latest = project / 'latest.json'
            latest.write_text(json.dumps(data))
            result = generate_compare_report(str(project), str(latest), None, 'partial.xlsx', logger_func=lambda _: None)
            report = load_workbook(result)
            sheet = report['実行結果']
            self.assertEqual(sheet['B3'].value, '一部操作を確認できず')
            self.assertIn('未確認', sheet['G3'].value)
            self.assertEqual(sheet['F3'].data_type, 's')
            report.close()

    def test_toggle_disables_baseline_and_switches_output_name(self):
        app = QApplication.instance() or QApplication([])
        ui = Ui_MainWindow()
        tab = ui._create_analysis_tab()
        owner = SimpleNamespace(ui=ui)
        ui.txt_f1.setText('/tmp/latest.json')
        ui.txt_f2.setText('/tmp/baseline.json')
        self.assertFalse(ui.chk_compare.isChecked())
        self.assertFalse(ui.comparison_input.isHidden())
        AutoTrackerApp.update_comparison_output_name(owner)
        self.assertFalse(ui.comparison_input.isHidden())
        self.assertFalse(ui.txt_f2.isEnabled())
        self.assertFalse(ui.btn_f2.isEnabled())
        self.assertEqual(ui.txt_out.text(), 'latest.xlsx')
        self.assertEqual(ui.lbl_analysis_input.text(), '計測データ')
        ui.chk_compare.setChecked(True)
        self.assertFalse(ui.comparison_input.isHidden())
        self.assertTrue(ui.txt_f2.isEnabled())
        self.assertTrue(ui.btn_f2.isEnabled())
        self.assertEqual(ui.txt_f2.text(), '/tmp/baseline.json')
        self.assertEqual(ui.btn_run_analysis.text(), 'レポートを生成')
        tab.close()


if __name__ == '__main__':
    unittest.main()
