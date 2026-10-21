import unittest

from openpyxl import Workbook
from excel_reporter import load_project_config, process_single_sheet
from measurement_adapters.adobe_analytics import AdobeAnalyticsAdapter, legacy_request_data


class AppMeasurementPostTests(unittest.TestCase):
    def test_capture_preserves_body_and_merges_query(self):
        body = 'events=event1%2Cevent2&v1=%E6%97%A5%E6%9C%AC+%2B&empty=&dup=a&dup=b'
        packet = legacy_request_data('https://example.com/b/ss/suite/1?events=old&pageName=test',
                                     'POST', body)
        self.assertEqual(packet['post_data'], body)
        self.assertEqual(packet['params'], {'events': 'event1,event2', 'pageName': 'test',
                                          'v1': '日本 +', 'empty': '', 'dup': ['a', 'b']})

    def test_excel_outputs_body_fields_and_discovers_unknown_fields(self):
        packet = legacy_request_data('https://example.com/b/ss/suite/1?pageName=test',
                                     'POST', 'events=event1&v1=abc')
        packet.update(step_id='0001', step_index=0)
        wb = Workbook()
        ws = wb.active
        ws.title = 'Adobe Analytics'
        ws.append(['Adobe Analytics'])
        ws.append(['送信文字列', 'XDM', 'Data', '名前', '説明'])
        ws.append(['events', '', '', 'Events', ''])
        process_single_sheet(ws, ws.title, [packet], [], load_project_config('/nonexistent'),
                             'latest.json', '', compare=False)
        self.assertEqual(ws['F3'].value, 'event1')
        rows = {ws.cell(r, 1).value: ws.cell(r, 6).value for r in range(3, ws.max_row + 1)}
        self.assertEqual(rows['v1'], 'abc')
        self.assertEqual(AdobeAnalyticsAdapter().packet_identity(packet)['events'], 'event1')
        wb.close()

    def test_legacy_get_and_old_json_remain_readable(self):
        adapter = AdobeAnalyticsAdapter()
        packet = legacy_request_data('https://example.com/b/ss/suite/1?events=event2', 'GET')
        self.assertEqual(adapter.extract_value(packet, {'source_key': 'events'}), 'event2')
        self.assertEqual(adapter.extract_value({'params': {'events': 'event3'}},
                                               {'source_key': 'events'}), 'event3')


if __name__ == '__main__':
    unittest.main()
