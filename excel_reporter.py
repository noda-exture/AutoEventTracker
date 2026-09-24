import os
import sys
import json
from datetime import datetime
from functools import lru_cache
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from measurement_adapters import get_adapter, get_registered_adapters
from measurement_adapters.ga4 import parse_query_string

def load_project_config(project_dir):
    config_path = os.path.join(project_dir, "config.json")
    default_config = {
        "client_name": "クライアント名未設定",
        "task_name": "課題名未設定",
        "excel_setting": {
            "font_name": "游ゴシック",
            "theme_color_new": "E2EFDA",
            "theme_color_old": "FFF2CC",
            "diff_color": "FFC7CE",
            "diff_font_color": "9C0006"
        },
        "extraction_rules": {
            "ignore_params": ["_ts", "gtm_auth", "gjid", "gtm", "requestId", "configId"],
            "nested_separator": "."
        }
    }
    
    if os.path.exists(config_path):
        with open(config_path, "r", encoding="utf-8") as f:
            user_config = json.load(f)
            if "excel_setting" in user_config:
                default_config["excel_setting"].update(user_config["excel_setting"])
            if "extraction_rules" in user_config:
                default_config["extraction_rules"].update(user_config["extraction_rules"])
            default_config.update({k: v for k, v in user_config.items() if k not in ["excel_setting", "extraction_rules"]})
            
    return default_config

def replace_placeholders_in_sheet(ws, config):
    today_str = datetime.now().strftime("%Y年%m月%d日")
    placeholders = {
        "{client_name}": config.get("client_name", ""),
        "{task_name}": config.get("task_name", ""),
        "{date}": today_str
    }
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val_str = cell.value
                for placeholder, actual_value in placeholders.items():
                    if placeholder in val_str:
                        val_str = val_str.replace(placeholder, actual_value)
                cell.value = val_str

def parse_ga4_query_string(raw_str):
    """後方互換用。GA4固有の解析はGA4Adapterへ移管済み。"""
    return parse_query_string(raw_str)

def get_nested_value(d, path, separator="."):
    """後方互換用。Adobe固有の解析はAdobeAnalyticsAdapterへ移管済み。"""
    from measurement_adapters.adobe_analytics import get_nested_value as adobe_get_nested_value

    return adobe_get_nested_value(d, path, separator)

def search_value_by_mappings(packet, m_row):
    """後方互換用。AdobeAnalyticsAdapterの抽出処理を呼び出す。"""
    return get_adapter("AA").extract_value(packet, m_row)

def format_extracted_value(val):
    if val is None: return ""
    if isinstance(val, (dict, list)): return json.dumps(val, ensure_ascii=False)
    return str(val)

def filter_events_by_type(events, target_type):
    """後方互換用。登録済みアダプターへイベント判定を委譲する。"""
    return get_adapter(target_type).filter_events(events)


def packet_identity(packet, target_type):
    """後方互換用。計測サービス固有の識別処理をアダプターへ委譲する。"""
    return get_adapter(target_type).packet_identity(packet)


def packet_match_score(packet_new, packet_old, target_type):
    new_identity = packet_identity(packet_new, target_type)
    old_identity = packet_identity(packet_old, target_type)
    if not new_identity or not old_identity or new_identity.get("tool") != old_identity.get("tool"):
        return -1000

    score = 20
    compared_semantic_field = False
    semantic_fields = ("event", "page", "url", "events")
    new_semantic_fields = {field for field in semantic_fields if new_identity.get(field)}
    old_semantic_fields = {field for field in semantic_fields if old_identity.get(field)}
    if (new_semantic_fields or old_semantic_fields) and not (new_semantic_fields & old_semantic_fields):
        return -1000

    new_event = new_identity.get("event")
    old_event = old_identity.get("event")
    if new_event and old_event:
        compared_semantic_field = True
        if new_event != old_event:
            return -1000
        score += 70

    for field, weight in (("page", 30), ("url", 25), ("events", 20), ("account", 10)):
        new_value = new_identity.get(field)
        old_value = old_identity.get(field)
        if new_value and old_value:
            compared_semantic_field = True
            if new_value == old_value:
                score += weight
            elif field in ("page", "url") and not (new_event and old_event):
                score -= 20

    if new_identity.get("source") == old_identity.get("source"):
        score += 5

    new_seq = int(packet_new.get("packet_index_in_step", 0) or 0)
    old_seq = int(packet_old.get("packet_index_in_step", 0) or 0)
    score += max(0, 20 - abs(new_seq - old_seq) * 3)

    # Older result files do not have enough identity fields. In that case,
    # retain order matching as a compatibility fallback within the step only.
    if not compared_semantic_field:
        score += 10
    return score


def _best_packet_matches(packets_new, packets_old, target_type, threshold=35):
    n_new, n_old = len(packets_new), len(packets_old)
    if not n_new or not n_old:
        return []

    scores = [[packet_match_score(new, old, target_type) for old in packets_old] for new in packets_new]

    if n_new <= 12 and n_old <= 12:
        @lru_cache(maxsize=None)
        def solve(new_index, used_old_mask):
            if new_index >= n_new:
                return 0, ()

            best_score, best_pairs = solve(new_index + 1, used_old_mask)
            for old_index in range(n_old):
                if used_old_mask & (1 << old_index):
                    continue
                match_score = scores[new_index][old_index]
                if match_score < threshold:
                    continue
                tail_score, tail_pairs = solve(new_index + 1, used_old_mask | (1 << old_index))
                candidate_score = match_score + tail_score
                if candidate_score > best_score:
                    best_score = candidate_score
                    best_pairs = ((new_index, old_index, match_score),) + tail_pairs
            return best_score, best_pairs

        return list(solve(0, 0)[1])

    candidates = []
    for new_index in range(n_new):
        for old_index in range(n_old):
            if scores[new_index][old_index] >= threshold:
                candidates.append((scores[new_index][old_index], new_index, old_index))
    candidates.sort(reverse=True)
    used_new, used_old, matches = set(), set(), []
    for score, new_index, old_index in candidates:
        if new_index in used_new or old_index in used_old:
            continue
        used_new.add(new_index)
        used_old.add(old_index)
        matches.append((new_index, old_index, score))
    return matches


def match_packets_in_step(packets_new, packets_old, target_type):
    matches = _best_packet_matches(packets_new, packets_old, target_type)
    matched_new = {new_index for new_index, _, _ in matches}
    matched_old = {old_index for _, old_index, _ in matches}
    pairs = [(packets_new[new_index], packets_old[old_index], score) for new_index, old_index, score in matches]
    pairs.extend((packet, None, None) for index, packet in enumerate(packets_new) if index not in matched_new)
    pairs.extend((None, packet, None) for index, packet in enumerate(packets_old) if index not in matched_old)

    def pair_order(pair):
        new_packet, old_packet, _ = pair
        packet = new_packet or old_packet or {}
        return (
            int(packet.get("packet_index_in_step", 0) or 0),
            float(packet.get("timestamp", 0) or 0),
        )

    return sorted(pairs, key=pair_order)


def _group_events(events):
    groups = {}
    for event in events:
        step_index = int(event.get("step_index", 0) or 0)
        step_id = str(event.get("step_id") or f"legacy_{step_index + 1:04d}")
        if step_id.startswith("step_") and len(step_id) > len("step_"):
            step_id = step_id[len("step_"):]
        if step_id not in groups:
            groups[step_id] = {
                "step_id": step_id,
                "step_index": step_index,
                "step_name": event.get("step_name") or "",
                "events": [],
            }
        if "packet_index_in_step" not in event:
            event["packet_index_in_step"] = len(groups[step_id]["events"])
        groups[step_id]["events"].append(event)
    return groups


def align_step_groups(events_new, events_old):
    """Align by stable step ID first, then by index for legacy result files."""
    new_groups = _group_events(events_new)
    old_groups = _group_events(events_old)
    aligned = []
    used_old = set()

    for new_id, new_group in new_groups.items():
        old_group = old_groups.get(new_id)
        if old_group:
            used_old.add(new_id)
        else:
            old_group = next(
                (group for old_id, group in old_groups.items()
                 if old_id not in used_old and group["step_index"] == new_group["step_index"]),
                None,
            )
            if old_group:
                used_old.add(old_group["step_id"])
        aligned.append({
            "step_id": new_group["step_id"],
            "step_index": new_group["step_index"],
            "step_name": new_group["step_name"] or (old_group or {}).get("step_name", ""),
            "new": new_group["events"],
            "old": old_group["events"] if old_group else [],
        })

    for old_id, old_group in old_groups.items():
        if old_id in used_old:
            continue
        aligned.append({
            "step_id": old_group["step_id"],
            "step_index": old_group["step_index"],
            "step_name": old_group["step_name"],
            "new": [],
            "old": old_group["events"],
        })

    return sorted(aligned, key=lambda group: (group["step_index"], group["step_id"]))


def packet_display_name(packet, target_type):
    return get_adapter(target_type).packet_display_name(packet)

def process_single_sheet(ws, sheet_name, ev_new, ev_old, config, latest_filename, base_filename):
    ex_set = config["excel_setting"]
    ext_rules = config["extraction_rules"]
    replace_placeholders_in_sheet(ws, config)
    adapter = get_adapter(sheet_name)
    
    START_DATA_COL = 1
    while ws.cell(row=2, column=START_DATA_COL).value:
        START_DATA_COL += 1

    mapping_rows = []
    url_row_idx = None
    for r in range(3, ws.max_row + 1):
        if START_DATA_COL == 4:
            source_key = str(ws.cell(row=r, column=1).value or "").strip()
            label = str(ws.cell(row=r, column=2).value or "").strip()
            desc = str(ws.cell(row=r, column=3).value or "").strip()
            primary_path, secondary_path = "", ""
        else:
            source_key = str(ws.cell(row=r, column=1).value or "").strip()
            primary_path = str(ws.cell(row=r, column=2).value or "").strip()
            secondary_path = str(ws.cell(row=r, column=3).value or "").strip()
            label = str(ws.cell(row=r, column=4).value or "").strip()
            desc = str(ws.cell(row=r, column=5).value or "").strip()
            
        if source_key or label:
            mapping_rows.append({
                "row_idx": r,
                "source_key": source_key,
                "primary_path": primary_path,
                "secondary_path": secondary_path,
                "label": label,
                "description": desc,
            })
            if not url_row_idx and (source_key in ["g", "dl"] or "url" in label.lower()):
                url_row_idx = r

    OLD_START_ROW = 2 + len(mapping_rows) + 4
    ws.row_dimensions[2].height = 48
    ws.row_dimensions[OLD_START_ROW].height = 48
    
    fills = {
        "new": PatternFill(start_color=ex_set["theme_color_new"], end_color=ex_set["theme_color_new"], fill_type="solid"),
        "old": PatternFill(start_color=ex_set["theme_color_old"], end_color=ex_set["theme_color_old"], fill_type="solid"),
        "diff": PatternFill(start_color=ex_set["diff_color"], end_color=ex_set["diff_color"], fill_type="solid"),
        "empty": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    }
    fonts = {
        "diff": Font(name=ex_set["font_name"], bold=True, color=ex_set["diff_font_color"], size=9),
        "normal": Font(name=ex_set["font_name"], size=9),
        "header": Font(name=ex_set["font_name"], bold=True, size=10)
    }
    thin_border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                         top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

    TITLE_COL = 4
    ws.cell(row=1, column=TITLE_COL, value=f"【 現在データ (新) 】 ファイル: {latest_filename}").font = Font(name=ex_set["font_name"], bold=True, size=11, color="2E6930")
    ws.cell(row=OLD_START_ROW - 1, column=TITLE_COL, value=f"【 過去データ (旧) 】 ファイル: {base_filename}").font = Font(name=ex_set["font_name"], bold=True, size=11, color="4B6F96")

    for idx, m_row in enumerate(mapping_rows):
        target_old_row = OLD_START_ROW + 1 + idx
        for col_i in range(1, START_DATA_COL):
            orig_val = ws.cell(row=m_row["row_idx"], column=col_i).value
            c_target = ws.cell(row=target_old_row, column=col_i, value=orig_val)
            c_target.font = fonts["normal"]
            c_target.fill = PatternFill(start_color="FAFAF7", end_color="FAFAF7", fill_type="solid")
            c_target.border = thin_border

    current_col = START_DATA_COL
    total_output_cols = 0
    aligned_steps = align_step_groups(ev_new, ev_old)

    for step_group in aligned_steps:
        packet_pairs = match_packets_in_step(step_group["new"], step_group["old"], adapter)
        if not packet_pairs:
            continue

        step_index = step_group["step_index"]
        step_label = "Start" if step_index < 0 else f"Step {step_index + 1}"

        for p_idx, (p_new, p_old, _match_score) in enumerate(packet_pairs):
            display_packet = p_new or p_old
            packet_name = adapter.packet_display_name(display_packet)
            p_label = f"{step_label}-{p_idx + 1}\n{packet_name}"
            if p_new and not p_old:
                p_label += "\n[最新のみ]"
            elif p_old and not p_new:
                p_label += "\n[比較元のみ]"
                
            for cell_h, fill_h in [(ws.cell(row=2, column=current_col, value=p_label), fills["new"]),
                                   (ws.cell(row=OLD_START_ROW, column=current_col, value=p_label), fills["old"])]:
                cell_h.font = fonts["header"]
                cell_h.fill = fill_h
                cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell_h.border = thin_border

            for idx, m_row in enumerate(mapping_rows):
                row_new = m_row["row_idx"]
                row_old = OLD_START_ROW + 1 + idx
                val_new = val_old = ""

                res_new = adapter.extract_value(p_new, m_row)
                res_old = adapter.extract_value(p_old, m_row)
                if res_new is not None:
                    val_new = format_extracted_value(res_new)
                if res_old is not None:
                    val_old = format_extracted_value(res_old)

                c_new = ws.cell(row=row_new, column=current_col, value=val_new)
                c_old = ws.cell(row=row_old, column=current_col, value=val_old)
                
                for c in [c_new, c_old]:
                    c.font = fonts["normal"]
                    c.border = thin_border
                    c.alignment = Alignment(horizontal="left", vertical="center")

                if not p_new and p_old:
                    c_new.value = "最新になし"
                elif p_new and not p_old:
                    c_old.value = "比較元になし"

            current_col += 1
            total_output_cols += 1

    if total_output_cols == 0: return

    # 比較結果の色はセルへ直接設定せず、値の変更に追従する条件付き書式で付ける。
    first_data_col = get_column_letter(START_DATA_COL)
    last_data_col = get_column_letter(START_DATA_COL + total_output_cols - 1)
    for idx, m_row in enumerate(mapping_rows):
        row_new = m_row["row_idx"]
        row_old = OLD_START_ROW + 1 + idx
        new_range = f"{first_data_col}{row_new}:{last_data_col}{row_new}"
        old_range = f"{first_data_col}{row_old}:{last_data_col}{row_old}"

        ws.conditional_formatting.add(
            new_range,
            FormulaRule(
                formula=[f'{first_data_col}{row_new}="最新になし"'],
                stopIfTrue=True,
                fill=fills["empty"],
            ),
        )
        ws.conditional_formatting.add(
            old_range,
            FormulaRule(
                formula=[f'{first_data_col}{row_old}="比較元になし"'],
                stopIfTrue=True,
                fill=fills["empty"],
            ),
        )

        if m_row["source_key"] not in ext_rules["ignore_params"]:
            ws.conditional_formatting.add(
                new_range,
                FormulaRule(
                    formula=[
                        f'AND({first_data_col}{row_new}<>"",'
                        f'{first_data_col}{row_new}<>{first_data_col}{row_old},'
                        f'{first_data_col}{row_new}<>"最新になし")'
                    ],
                    fill=fills["diff"],
                    font=fonts["diff"],
                ),
            )
            ws.conditional_formatting.add(
                old_range,
                FormulaRule(
                    formula=[
                        f'AND({first_data_col}{row_old}<>"",'
                        f'{first_data_col}{row_old}<>{first_data_col}{row_new},'
                        f'{first_data_col}{row_old}<>"比較元になし")'
                    ],
                    fill=fills["diff"],
                    font=fonts["diff"],
                ),
            )

    ws.row_dimensions.group(3, 2 + len(mapping_rows), hidden=False, outline_level=1)
    ws.row_dimensions.group(OLD_START_ROW + 1, OLD_START_ROW + len(mapping_rows), hidden=False, outline_level=1)

    for idx, m_row in enumerate(mapping_rows):
        row_new = m_row["row_idx"]
        row_old = OLD_START_ROW + 1 + idx
        is_empty_row = True
        
        for c_idx in range(total_output_cols):
            col_num = START_DATA_COL + c_idx
            v_new = str(ws.cell(row=row_new, column=col_num).value or "").strip()
            v_old = str(ws.cell(row=row_old, column=col_num).value or "").strip()
            
            if (v_new and v_new not in ["最新になし", "比較元になし"]) or \
               (v_old and v_old not in ["最新になし", "比較元になし"]):
                is_empty_row = False
                break
                
        if is_empty_row:
            ws.row_dimensions[row_new].hidden = True
            ws.row_dimensions[row_old].hidden = True

    if url_row_idx:
        start_col_grp = START_DATA_COL
        current_url = str(ws.cell(row=url_row_idx, column=START_DATA_COL).value or "").strip()
        
        for c_idx in range(1, total_output_cols):
            col_num = START_DATA_COL + c_idx
            url_val = str(ws.cell(row=url_row_idx, column=col_num).value or "").strip()
            
            if url_val != current_url:
                if col_num - 1 > start_col_grp:
                    for grp_col in range(start_col_grp + 1, col_num):
                        col_letter = get_column_letter(grp_col)
                        ws.column_dimensions[col_letter].outline_level = 1
                start_col_grp = col_num
                current_url = url_val
                
        if (START_DATA_COL + total_output_cols - 1) > start_col_grp:
            for grp_col in range(start_col_grp + 1, START_DATA_COL + total_output_cols):
                col_letter = get_column_letter(grp_col)
                ws.column_dimensions[col_letter].outline_level = 1

    ws.sheet_properties.outlinePr.summaryRight = False
    ws.sheet_properties.outlinePr.summaryBelow = False

    # A～C列だけを非表示にし、項目・説明のD/E列と比較データ列は表示する。
    max_col_to_check = START_DATA_COL + total_output_cols - 1
    for c_idx in range(1, max_col_to_check + 1):
        col_letter = get_column_letter(c_idx)
        if c_idx <= 3:
            ws.column_dimensions[col_letter].hidden = True
        else:
            ws.column_dimensions[col_letter].hidden = False
            
            # 列幅の自動調整もここで行う
            max_len = 0
            for cell in ws[col_letter]:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)


def generate_compare_report(
    project_name,
    latest_json_path,
    base_json_path,
    output_excel_name,
    logger_func=print,
    adapters=None,
):
    project_dir = os.path.join("project", project_name)
    template_path = os.path.join(project_dir, "template.xlsx")
    
    l_path = latest_json_path if os.path.isabs(latest_json_path) else os.path.join(project_dir, "outputs", latest_json_path)
    b_path = base_json_path if os.path.isabs(base_json_path) else os.path.join(project_dir, "outputs", base_json_path)
    output_excel_path = os.path.join(project_dir, "outputs", output_excel_name)

    config = load_project_config(project_dir)
    
    if not os.path.exists(template_path): return logger_func(f"テンプレートが見つかりません: {template_path}")
    if not os.path.exists(l_path): return logger_func(f"最新の計測JSONが見つかりません: {l_path}")
    if not os.path.exists(b_path): return logger_func(f"比較元の計測JSONが見つかりません: {b_path}")

    try:
        with open(l_path, "r", encoding="utf-8") as f: latest_data = json.load(f)
        with open(b_path, "r", encoding="utf-8") as f: base_data = json.load(f)
    except json.JSONDecodeError:
        return logger_func("JSONファイルの読み込みに失敗しました。ファイルが壊れていないか確認してください。")

    latest_events = latest_data.get("events", [])
    base_events = base_data.get("events", [])

    latest_scenario = latest_data.get("scenario")
    base_scenario = base_data.get("scenario")
    if latest_scenario and base_scenario and latest_scenario != base_scenario:
        logger_func(
            f"注意: 異なるシナリオを比較しています "
            f"(最新: {latest_scenario} / 比較元: {base_scenario})"
        )
    
    wb = load_workbook(template_path)

    # 登録済みアダプターごとに独立してイベントを抽出し、対応シートを処理する。
    process_targets = []
    configured_adapters = tuple(adapters) if adapters is not None else get_registered_adapters()
    active_adapters = tuple(get_adapter(adapter) for adapter in configured_adapters)
    for adapter in active_adapters:
        events_new = adapter.filter_events(latest_events)
        events_old = adapter.filter_events(base_events)
        if not events_new and not events_old:
            continue
        if adapter.sheet_name in wb.sheetnames:
            process_targets.append((adapter, events_new, events_old))
        else:
            logger_func(
                f"{adapter.display_name}のデータがありますが、"
                f"テンプレートに『{adapter.sheet_name}』シートが無いためスキップします。"
            )

    if not process_targets:
        supported_sheets = " / ".join(adapter.sheet_name for adapter in active_adapters)
        return logger_func(
            "処理対象のデータに対応するシートが見つかりません。"
            f"テンプレートのシート名（{supported_sheets}）を確認してください。"
        )

    keep_sheets = ["表紙"] + [target[0].sheet_name for target in process_targets]
    for name in list(wb.sheetnames):
        if name not in keep_sheets:
            wb.remove(wb[name])

    if "表紙" in wb.sheetnames:
        replace_placeholders_in_sheet(wb["表紙"], config)

    for adapter, ev_new, ev_old in process_targets:
        logger_func(f"シート『{adapter.sheet_name}』の比較処理を行っています。")
        process_single_sheet(
            ws=wb[adapter.sheet_name],
            sheet_name=adapter,
            ev_new=ev_new,
            ev_old=ev_old,
            config=config,
            latest_filename=os.path.basename(latest_json_path),
            base_filename=os.path.basename(base_json_path)
        )

    try:
        wb.save(output_excel_path)
    except PermissionError:
        logger_func(f"【エラー】出力先のExcelファイル ({output_excel_name}) が開かれています。\nExcelを閉じてから再度実行してください。\n")
        return None

    logger_func(f"比較レポートを生成しました。\n保存先: {output_excel_path}\n")
    return output_excel_path

if __name__ == "__main__":
    if len(sys.argv) > 4:
        generate_compare_report(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4])
    else:
        print("使い方: python excel_reporter.py [案件名] [最新JSON名] [過去JSON名] [出力Excel名]")
